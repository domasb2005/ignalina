import os
from docx import Document
from PyPDF2 import PdfReader, PdfWriter
from PyPDF2 import PageObject

def generate_pdf(receiver, freq):
    # Load the base DOCX file
    document = Document("./static/base.docx")

    # Replace placeholders in the document
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if "YYYYY" in run.text:
                run.text = run.text.replace("YYYYY", receiver)
                run.bold = True  # Make the replaced text bold
            if "yyyyy" in run.text:
                run.text = run.text.replace("yyyyy", receiver)
                run.bold = True  # Make the replaced text bold
            if "XXXX" in run.text:
                run.text = run.text.replace("XXXX", str(freq))

    # Save the modified document to a new DOCX file
    modified_docx = "./static/document.docx"
    document.save(modified_docx)

    # Convert DOCX to PDF using libreoffice
    os.system("libreoffice --headless --convert-to pdf {} --outdir ./static".format(modified_docx))

    # # Apply formatting to the converted PDF
    # input_pdf_path = "./static/document.pdf"
    # output_pdf_path = "./static/modified_document.pdf"

    # with open(input_pdf_path, "rb") as input_pdf_file:
    #     reader = PdfReader(input_pdf_file)
    #     writer = PdfWriter()

    #     for page in reader.pages:
    #         page_text = page.extract_text()
    #         page_text = page_text.replace("YYYYY", f"<b>{receiver}</b>")
    #         page_text = page_text.replace("yyyyy", f"<b>{receiver}</b>")
    #         page_text = page_text.replace("XXXX", f"<b>{str(freq)}</b>")
            
    #         modified_page = PageObject.create_page(page_text)
    #         writer.add_page(modified_page)

    #     with open(output_pdf_path, "wb") as output_pdf_file:
    #         writer.write(output_pdf_file)

# Example usage
generate_pdf("test", 9999)
